from flask import Flask,request, render_template, flash, redirect, url_for,session, logging, send_file
from flask_mysqldb import MySQL 
from wtforms import Form, StringField, TextAreaField, PasswordField, validators, DateTimeField, BooleanField, IntegerField
from flask_wtf import FlaskForm
from flask_wtf.file import FileField, FileRequired, FileAllowed
from passlib.hash import sha256_crypt
from functools import wraps
from werkzeug.utils import secure_filename
from docx import Document
from coolname import generate_slug
from datetime import timedelta, datetime
from flask_mail import Mail, Message
from threading import Thread
from flask import render_template_string
from itsdangerous import URLSafeTimedSerializer
from validate_email import validate_email
import random
import json
import csv
import operator
from wtforms_components import TimeField
from wtforms.fields.html5 import DateField
from wtforms.validators import ValidationError
import socket
from emailverifier import Client

app = Flask(__name__)
app.secret_key= 'huihui'

#Config MySQL
app.config['MYSQL_HOST'] = 'localhost'
app.config['MYSQL_USER'] = ''
app.config['MYSQL_PASSWORD'] = ''
app.config['MYSQL_DB'] = 'flask'
app.config['MYSQL_CURSORCLASS'] = 'DictCursor'


app.config.update(
	DEBUG=True,
	MAIL_SERVER='smtp.gmail.com',
	MAIL_PORT=465,
	MAIL_USE_SSL=True,
	MAIL_USERNAME = '',
	MAIL_PASSWORD = ''
	)
mail = Mail(app)

def asynch(f):
	@wraps(f)
	def wrapper(*args, **kwargs):
		thr = Thread(target=f, args=args, kwargs=kwargs)
		thr.start()
	return wrapper

@asynch
def send_async_email(app, msg):
	with app.app_context():
		mail.send(msg)


htmlbody='''
Your account on <b>The Best</b> Quiz App was successfully created.
Please click the link below to confirm your email address and
activate your account:
  
<a href="{{ confirm_url }}">{{ confirm_url }}</a>
 <p>
--
Questions? Comments? Email </p>
'''

#email verifier api key
client = Client('at_rFxZz7zEX8CO8V5IDBfzexOe2fW8b')



def get_local_ip():	
	s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
	s.connect(('8.8.8.8', 1))
	local_ip_address = s.getsockname()[0]
	return local_ip_address


def send_email(recipients,html_body):
	try:
		msg = Message('Confirm Your Email Address',
		  sender="nickqwerty76@gmail.com",
		  recipients=recipients)
		msg.html = html_body
		send_async_email(app, msg)
		# return 'Mail sent!'
		return
	except Exception as e:
		# return(str(e))
		return


def send_confirmation_email(user_email):
	confirm_serializer = URLSafeTimedSerializer(app.config['SECRET_KEY'])
 
	confirm_url = url_for('confirm_email',
		token=confirm_serializer.dumps(user_email, salt='email-confirmation-salt'),
		_external=True)
	local_ip = get_local_ip()
	x=""
	if 'localhost' in confirm_url:
		x=confirm_url.split("localhost:5000")
	else:
		x=confirm_url.split("127.0.0.1:5000")
	confirm_url = x[0] + local_ip + ":5000"+ x[1]
	html = render_template_string(htmlbody, confirm_url=confirm_url)

	send_email([user_email], html)


# test this function whether correct link is produced##############

#init Mysql
mysql = MySQL(app)

@app.before_request
def make_session_permanent():
	session.permanent = True
	app.permanent_session_lifetime = timedelta(minutes=5)


def is_logged(f):
	@wraps(f)
	def wrap(*args, **kwargs):
		if 'logged_in' in session:
			return f(*args, **kwargs)
		else:
			flash('Unauthorized, Please login','danger')
			return redirect(url_for('login'))
	return wrap


def doctodict(filepath):
	document = Document(filepath)
	data={}
	count=1
	for table in document.tables:
		temp = {}
		for rowNo,_ in enumerate(table.rows):
			temp[table.cell(rowNo, 0).text]=table.cell(rowNo, 1).text
		data[count] = temp
		count+=1
 
	return data


class RegisterForm(Form):
	name = StringField('Name', [validators.Length(min=3, max=50)])
	username = StringField('Username', [validators.Length(min=4,max=25)])
	email = StringField('Email', [validators.Email()])
	password = PasswordField('Password', [
			validators.Regexp("^(?=.*[A-Za-z])(?=.*\d)[A-Za-z\d]{8,}$", message="Password should contain min 8 characters including 1 letter and 1 number."),
			validators.DataRequired(),
			validators.EqualTo('confirm', message="Password do not match")
		])
	confirm = PasswordField('Confirm Password')


class UploadForm(FlaskForm):
	subject = StringField('Subject')
	topic = StringField('Topic')
	doc = FileField('Docx Upload', validators=[FileRequired()])
	start_date = DateField('Start Date')
	start_time = TimeField('Start Time', default=datetime.utcnow()+timedelta(hours=5.5))
	end_date = DateField('End Date')
	end_time = TimeField('End Time', default=datetime.utcnow()+timedelta(hours=5.5))
	show_result = BooleanField('Show Result after completion')
	neg_mark = BooleanField('Enable negative marking')
	duration = IntegerField('Duration(in min)')
	password = StringField('Test Password', [validators.Length(min=3, max=6)])

	def validate_end_date(form, field):
		if field.data < form.start_date.data:
			raise ValidationError("End date must not be earlier than start date.")
	
	def validate_end_time(form, field):
		start_date_time = datetime.strptime(str(form.start_date.data) + " " + str(form.start_time.data),"%Y-%m-%d %H:%M:%S").strftime("%Y-%m-%d %H:%M")
		end_date_time = datetime.strptime(str(form.end_date.data) + " " + str(field.data),"%Y-%m-%d %H:%M:%S").strftime("%Y-%m-%d %H:%M")
		if start_date_time >= end_date_time:
			raise ValidationError("End date time must not be earlier/equal than start date time")
	
	def validate_start_date(form, field):
		if datetime.strptime(str(form.start_date.data) + " " + str(form.start_time.data),"%Y-%m-%d %H:%M:%S") < datetime.now():
			raise ValidationError("Start date and time must not be earlier than current")

class TestForm(Form):
	test_id = StringField('Test ID')
	password = PasswordField('Test Password')


@app.route('/')
def index():
	return render_template('index.html')

@app.route('/register', methods=['GET','POST'])
def register():
	form = RegisterForm(request.form)
	if request.method == 'POST' and form.validate():
		name = form.name.data 
		email = form.email.data

		#email verifier	
		data = client.get(email)
		if str(data.smtp_check) == 'False':
			flash('Invalid email, please provide a valid email address','danger')
			return render_template('register.html', form=form)

		send_confirmation_email(email)

		username = form.username.data
		password = sha256_crypt.encrypt(str(form.password.data))
		cur = mysql.connection.cursor()
		cur.execute('INSERT INTO users(username,name,email, password,confirmed) values(%s,%s,%s,%s,0)', (username,name, email, password))
		mysql.connection.commit()
		cur.close()
		flash('Thanks for registering!  Please check your email to confirm your email address.', 'success')
		return redirect(url_for('login')) 
		# change in login function to redirect to warning page

	return render_template('register.html', form=form)

	
@app.route('/login', methods=['GET','POST'])
def login():
	if request.method == 'POST':
		username = request.form['username']
		password_candidate = request.form['password']
		cur = mysql.connection.cursor()
		results = cur.execute('SELECT * from users where username = %s' , [username])
		if results > 0:
			data = cur.fetchone()
			password = data['password']
			confirmed = data['confirmed']
			name = data['name']
			if confirmed == 0:
				error = 'Please confirm email before logging in'
				return render_template('login.html', error=error)
			if sha256_crypt.verify(password_candidate, password):
				session['logged_in'] = True
				session['username'] = username
				session['name'] = name
				return redirect(url_for('dashboard'))
			else:
				error = 'Invalid password'
				return render_template('login.html', error=error)
			cur.close()
		else:
			error = 'Username not found'
			return render_template('login.html', error=error)
	return render_template('login.html')


@app.route('/dashboard')
@is_logged
def dashboard():
	return render_template('dashboard.html')


@app.route('/logout')
def logout():
	session.clear()
	flash('Successfully logged out', 'success')
	return redirect(url_for('index'))


@app.route('/create-test', methods = ['GET', 'POST'])
@is_logged
def create_test():
	form = UploadForm()
	if request.method == 'POST' and form.validate_on_submit():
		f = form.doc.data
		filename = secure_filename(f.filename)
		f.save('questions/' + filename)
		cur = mysql.connection.cursor()
		d = doctodict('questions/' + f.filename.replace(' ', '_').replace('(','').replace(')',''))
		test_id = generate_slug(2)
		try:
			for no, data in d.items():
				marks = data['((MARKS)) (1/2/3...)']
				a = data['((OPTION_A))']
				b = data['((OPTION_B))']
				c = data['((OPTION_C))']
				d = data['((OPTION_D))']
				question = data['((QUESTION))']
				correct_ans = data['((CORRECT_CHOICE)) (A/B/C/D)']
				explanation = data['((EXPLANATION)) (OPTIONAL)']
				
				cur.execute('INSERT INTO questions(test_id,qid,q,a,b,c,d,ans,marks,explanation) values(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)', 
					(test_id,no,question,a,b,c,d,correct_ans,marks,explanation))
				mysql.connection.commit()
			start_date = form.start_date.data
			end_date = form.end_date.data
			start_time = form.start_time.data
			end_time = form.end_time.data
			start_date_time = str(start_date) + " " + str(start_time)
			end_date_time = str(end_date) + " " + str(end_time)
			show_result = form.show_result.data
			neg_mark = form.neg_mark.data

			duration = int(form.duration.data)*60
			password = form.password.data
			subject = form.subject.data
			topic = form.topic.data
			cur.execute('INSERT INTO teachers (username, test_id, start, end, duration, show_ans, password, subject, topic,neg_mark) values(%s,%s,%s,%s,%s,%s,%s, %s,%s,%s)',
				(dict(session)['username'], test_id, start_date_time, end_date_time, duration, show_result, password, subject, topic, neg_mark))
			mysql.connection.commit()
			cur.close()
			flash(f'Test ID: {test_id}', 'success')
			return redirect(url_for('dashboard'))
		except Exception as e:
			flash('Invalid Input File Format','danger')
			return redirect(url_for('create_test'))
		
	return render_template('create_test.html' , form = form)


@app.route('/give-test/<testid>', methods=['GET','POST'])
@is_logged
def test(testid):
	global duration,marked_ans
	if request.method == 'GET':
		try:
			data = {'duration': duration, 'marks': '', 'q': '', 'a': "", 'b':"",'c':"",'d':"" }
			return render_template('quiz.html' ,**data, answers=marked_ans)
		except:
			return redirect(url_for('give_test'))
	else:
		cur = mysql.connection.cursor()
		flag = request.form['flag']
		if flag == 'get':
			num = request.form['no']
			results = cur.execute('SELECT * from questions where test_id = %s and qid =%s',(testid, num))
			if results > 0:
				data = cur.fetchone()
				del data['ans']
				cur.close()
				return json.dumps(data)
		elif flag=='mark':
			qid = request.form['qid']
			ans = request.form['ans']
			results = cur.execute('SELECT * from students where test_id =%s and qid = %s and username = %s', (testid, qid, session['username']))
			if results > 0:
				cur.execute('UPDATE students set ans = %s where test_id = %s and qid = %s and username = %s', (testid, qid, session['username']))
			else:
				cur.execute('INSERT INTO students values(%s,%s,%s,%s)', (session['username'], testid, qid, ans))
			mysql.connection.commit()
			cur.close()
		elif flag=='time':
			time_left = request.form['time']
			try:
				cur.execute('UPDATE studentTestInfo set time_left=SEC_TO_TIME(%s) where test_id = %s and username = %s and completed=0', (time_left, testid, session['username']))
				mysql.connection.commit()
				cur.close()
			except:
				pass
		else:			
			cur.execute('UPDATE studentTestInfo set completed=1,time_left=sec_to_time(0) where test_id = %s and username = %s', (testid, session['username']))
			mysql.connection.commit()
			cur.close()
			flash("Test submitted successfully", 'info')
			return json.dumps({'sql':'fired'})


@app.route("/give-test", methods = ['GET', 'POST'])
@is_logged
def give_test():
	global duration, marked_ans	
	form = TestForm(request.form)
	if request.method == 'POST' and form.validate():
		test_id = form.test_id.data
		password_candidate = form.password.data
		cur = mysql.connection.cursor()
		results = cur.execute('SELECT * from teachers where test_id = %s', [test_id])
		if results > 0:
			data = cur.fetchone()
			password = data['password']
			duration = data['duration']
			start = data['start']
			start = str(start)
			end = data['end']
			end = str(end)
			if password == password_candidate:
				now = datetime.now()
				now = now.strftime("%Y-%m-%d %H:%M:%S")
				now = datetime.strptime(now,"%Y-%m-%d %H:%M:%S")
				if datetime.strptime(start,"%Y-%m-%d %H:%M:%S") < now and datetime.strptime(end,"%Y-%m-%d %H:%M:%S") > now:
					results = cur.execute('SELECT time_to_sec(time_left) as time_left,completed from studentTestInfo where username = %s and test_id = %s', (session['username'], test_id))
					if results > 0:
						results = cur.fetchone()
						is_completed = results['completed']
						if is_completed == 0:
							time_left = results['time_left']
							if time_left <= duration:
								duration = time_left
								results = cur.execute('SELECT * from students where username = %s and test_id = %s', (session['username'], test_id))
								marked_ans = {}
								if results > 0:
									results = cur.fetchall()
									for row in results:
										marked_ans[row['qid']] = row['ans']
									marked_ans = json.dumps(marked_ans)
									#return redirect(url_for('test' , testid = test_id))
						else:
							flash('Test already given', 'success')
							return redirect(url_for('give_test'))
					else:
						cur.execute('INSERT into studentTestInfo (username, test_id,time_left) values(%s,%s,SEC_TO_TIME(%s))', (session['username'], test_id, duration))
						mysql.connection.commit()
						results = cur.execute('SELECT time_to_sec(time_left) as time_left,completed from studentTestInfo where username = %s and test_id = %s', (session['username'], test_id))
						if results > 0:
							results = cur.fetchone()
							is_completed = results['completed']
							if is_completed == 0:
								time_left = results['time_left']
								if time_left <= duration:
									duration = time_left
									results = cur.execute('SELECT * from students where username = %s and test_id = %s', (session['username'], test_id))
									marked_ans = {}
									if results > 0:
										results = cur.fetchall()
										for row in results:
											marked_ans[row['qid']] = row['ans']
										marked_ans = json.dumps(marked_ans)
				else:
					if datetime.strptime(start,"%Y-%m-%d %H:%M:%S") > now:
						flash(f'Test start time is {start}', 'danger')
					else:
						flash(f'Test has ended', 'danger')
					return redirect(url_for('give_test'))
				return redirect(url_for('test' , testid = test_id))
			else:
				flash('Invalid password', 'danger')
				return redirect(url_for('give_test'))
		flash('Invalid testid', 'danger')
		return redirect(url_for('give_test'))
		cur.close()
	return render_template('give_test.html', form = form)


@app.route('/randomize', methods = ['POST'])
def random_gen():
	if request.method == "POST":
		id = request.form['id']
		cur = mysql.connection.cursor()
		results = cur.execute('SELECT count(*) from questions where test_id = %s', [id])
		if results > 0:
			data = cur.fetchone()
			total = data['count(*)']
			nos = list(range(1,int(total)+1))
			random.Random(id).shuffle(nos)
			cur.close()
			return json.dumps(nos)


@app.route('/<username>/<testid>')
@is_logged
def check_result(username, testid):
	if username == session['username']:
		cur = mysql.connection.cursor()
		results = cur.execute('SELECT * FROM teachers where test_id = %s', [testid])
		if results>0:
			results = cur.fetchone()
			check = results['show_ans']
			if check == 1:
				results = cur.execute('select explanation,q,a,b,c,d,marks,q.qid as qid, \
					q.ans as correct, ifnull(s.ans,0) as marked from questions q left join \
					students s on  s.test_id = q.test_id and s.test_id = %s \
					and s.username = %s and s.qid = q.qid group by q.qid \
					order by LPAD(lower(q.qid),10,0) asc', (testid, username))
				if results > 0:
					results = cur.fetchall()
					return render_template('tests_result.html', results= results)
			else:
				flash('You are not authorized to check the result', 'danger')
				return redirect(url_for('tests_given',username = username))
	else:
		return redirect(url_for('dashboard'))

#tests==dict in tuple

def neg_marks(username,testid):
	cur=mysql.connection.cursor()
	results = cur.execute("select marks,q.qid as qid, \
				q.ans as correct, ifnull(s.ans,0) as marked from questions q left join \
				students s on  s.test_id = q.test_id and s.test_id = %s \
				and s.username = %s and s.qid = q.qid group by q.qid \
				order by LPAD(lower(q.qid),10,0) asc", (testid, username))
	data=cur.fetchall()

	sum=0.0
	for i in range(results):
		if(str(data[i]['marked']).upper() != '0'):
			if(str(data[i]['marked']).upper() != str(data[i]['correct'])):
				sum=sum-0.25*int(data[i]['marks'])
			elif(str(data[i]['marked']).upper() == str(data[i]['correct'])):
				sum+=int(data[i]['marks'])
	return sum

def totmarks(username,tests): 
	cur = mysql.connection.cursor()
	for test in tests:
		testid = test['test_id']
		results=cur.execute("select neg_mark from teachers where test_id=%s",[testid])
		results=cur.fetchone()
		if results['neg_mark']==1:
			test['marks'] = neg_marks(username,testid) 

		else:
			results = cur.execute("select sum(marks) as totalmks from students s,questions q \
				where s.username=%s and s.test_id=%s and s.qid=q.qid and s.test_id=q.test_id \
				and s.ans=q.ans", (username, testid))	
			
			results = cur.fetchone()
			if str(results['totalmks']) == 'None':
				results['totalmks'] = 0
			test['marks'] = results['totalmks']
	return tests


def marks_calc(username,testid):
	if username == session['username']:
		cur = mysql.connection.cursor()
		results=cur.execute("select neg_mark from teachers where test_id=%s",[testid])
		results=cur.fetchone()
		if results['neg_mark']==1:
			return neg_marks(username,testid) 
		else:
			results = cur.execute("select sum(marks) as totalmks from students s,questions q where s.username=%s and s.test_id=%s and s.qid=q.qid and s.test_id=q.test_id and s.ans=q.ans", (username, testid))
			results = cur.fetchone()
			if str(results['totalmks']) == 'None':
				results['totalmks'] = 0
			return results['totalmks']

		
@app.route('/<username>/tests-given')
@is_logged
def tests_given(username):
	if username == session['username']:
		cur = mysql.connection.cursor()
		results = cur.execute('select distinct(students.test_id),subject,topic from students,teachers where students.username = %s and students.test_id=teachers.test_id', [username])
		results = cur.fetchall()
		results=totmarks(username,results)
		return render_template('tests_given.html', tests=results)
	else:
		flash('You are not authorized', 'danger')
		return redirect(url_for('dashboard'))

@app.route('/<username>/tests-created/<testid>', methods = ['POST','GET'])
@is_logged
def student_results(username, testid):
	if username == session['username']:
		cur = mysql.connection.cursor()
		results = cur.execute('select users.name as name,users.username as username,test_id from studentTestInfo,users where test_id = %s and completed = 1 and studentTestInfo.username=users.username ', [testid])
		results = cur.fetchall()
		final = []
		count = 1
		for user in results:
			score = marks_calc(user['username'], testid)
			user['srno'] = count
			user['marks'] = score
			final.append([count, user['name'], score])
			count+=1
		if request.method =='GET':
			results = sorted(results, key=operator.itemgetter('marks'))
			return render_template('student_results.html', data=results)
		else:
			fields = ['Sr No', 'Name', 'Marks']
			with open('static/' + testid + '.csv', 'w') as f:
				writer = csv.writer(f)
				writer.writerow(fields)
				writer.writerows(final)
			#return send_file('/static/' + testid + '.csv', as_attachment=True)

@app.route('/<username>/tests-created/<testid>/questions', methods = ['POST','GET'])
@is_logged
def questions(username, testid):
	if username == session['username']:
		cur = mysql.connection.cursor()
		results = cur.execute('SELECT * FROM teachers where test_id = %s', [testid])
		if results>0:
			results = cur.fetchone()
			results = cur.execute('select explanation,q,a,b,c,d,marks,q.qid as qid, \
				q.ans as correct, ifnull(s.ans,0) as marked from questions q left join \
				students s on  s.test_id = q.test_id and s.test_id = %s \
				and s.username = %s and s.qid = q.qid group by q.qid \
				order by LPAD(lower(q.qid),10,0) asc', (testid, username))
			if results > 0:
				results = cur.fetchall()
				return render_template('disp_questions.html', results= results)


@app.route('/<username>/tests-created')
@is_logged
def tests_created(username):
	if username == session['username']:
		cur = mysql.connection.cursor()
		results = cur.execute('select * from teachers where username = %s', [username])
		results = cur.fetchall()
		return render_template('tests_created.html', tests=results)
	else:
		flash('You are not authorized', 'danger')
		return redirect(url_for('dashboard'))


@app.route('/confirm/<token>')
def confirm_email(token):
	try:
		confirm_serializer = URLSafeTimedSerializer(app.config['SECRET_KEY'])
		email = confirm_serializer.loads(token, salt='email-confirmation-salt', max_age=3600)
	except:
		flash('The confirmation link is invalid or has expired.', 'error')
		return redirect(url_for('login'))

	cur = mysql.connection.cursor()
	results = cur.execute('SELECT * from users where email = %s' , [email])
	if results > 0:
		data = cur.fetchone()
		email_confirmed = data['confirmed']
		if email_confirmed:
			flash('Account already confirmed. Please login.', 'info')
		else:
			results = cur.execute('UPDATE users SET confirmed = 1 where email = %s' , [email])
			mysql.connection.commit()
			cur.close()
			flash('Thank you for confirming your email address!', 'success')
			return redirect(url_for('login'))
		return redirect(url_for('index'))


if __name__ == "__main__":
	app.run(host = "0.0.0.0",debug=True)
